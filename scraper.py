import json
import os
import time

from dotenv import load_dotenv
from google import genai
from google.genai import types
from google.genai.errors import APIError
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
from playwright.sync_api import sync_playwright
from pydantic import BaseModel

load_dotenv()


class UrunVerisi(BaseModel):
  urun_adi: str
  fiyat: str
  stokta_mi: bool


client = genai.Client(api_key=os.environ.get("GEMINI_API_KEY"))


def fetch_page_content(url: str) -> str:
  print(f"\n[{url}] adresine bağlanılıyor...")
  with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    context = browser.new_context(
        user_agent=(
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        ),
        viewport={"width": 1280, "height": 800},
    )
    page = context.new_page()

    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    page.mouse.wheel(0, 500)
    time.sleep(3)

    html_content = page.content()
    browser.close()
    return html_content


def analyze_with_gemini(html_content: str):
  prompt = (
      "Aşağıdaki HTML içeriğinden ürün/öğe bilgilerini ayıkla. Her bir öğe için"
      " tam adını, fiyatını (varsa, yoksa 'Belirtilmedi') ve stok durumunu"
      f" belirt:\n\nHTML:\n{html_content[:30000]}"
  )

  max_retries = 5
  for attempt in range(max_retries):
    try:
      response = client.models.generate_content(
          model="gemini-3.6-flash",
          contents=prompt,
          config=types.GenerateContentConfig(
              response_mime_type="application/json",
              response_schema=list[UrunVerisi],
              temperature=0.1,
          ),
      )
      return response.text
    except APIError as e:
      error_str = str(e)
      if (
          "429" in error_str
          or "503" in error_str
          or "UNAVAILABLE" in error_str
      ):
        wait_time = (attempt + 1) * 5
        print(
            f"Sunucu yoğun veya kota sınırı var ({attempt + 1}/{max_retries})."
            f" {wait_time} saniye beklenip tekrar deneniyor..."
        )
        time.sleep(wait_time)
      else:
        raise e
  raise Exception(
      "Sunucu yoğunluğu nedeniyle maksimum deneme sayısına ulaşıldı. Lütfen az"
      " sonra tekrar deneyin."
  )


def scrape_smart(url: str):
  raw_html = fetch_page_content(url)
  print("Sayfa içeriği alındı, Gemini API ile analiz ediliyor...")
  return analyze_with_gemini(raw_html)


def json_to_excel(json_filepath="output.json", excel_filepath="output.xlsx"):
  try:
    if not os.path.exists(json_filepath):
      print(f"⚠️ {json_filepath} dosyası bulunamadı.")
      return

    with open(json_filepath, "r", encoding="utf-8") as f:
      data = json.load(f)

    if not data:
      print("Dönüştürülecek veri bulunamadı.")
      return

    df = pd.DataFrame(data)

    with pd.ExcelWriter(excel_filepath, engine="openpyxl") as writer:
      df.to_excel(writer, index=False, sheet_name="Ürün Listesi")
      worksheet = writer.sheets["Ürün Listesi"]

      header_fill = PatternFill(
          start_color="1F4E78", end_color="1F4E78", fill_type="solid"
      )
      header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

      for col_num in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

      for col in worksheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    print(
        f"📊 Veriler başarıyla okunabilir Excel formatına aktarıldı:"
        f" {excel_filepath}"
    )

  except Exception as e:
    print(f"Excel dönüşüm hatası: {e}")


if __name__ == "__main__":
  user_url = input(
      "Kazımak istediğiniz Web Sayfası URL'sini yapıştırın: "
  ).strip()

  if not user_url:
    print("Geçerli bir URL girmediniz!")
  else:
    try:
      raw_result = scrape_smart(user_url)

      parsed_json = json.loads(raw_result)
      formatted_json = json.dumps(parsed_json, ensure_ascii=False, indent=2)

      print("\n--- ÇIKARILAN DÜZENLİ JSON VERİSİ ---")
      print(formatted_json)

      with open("output.json", "w", encoding="utf-8") as f:
        f.write(formatted_json)

      print("\nSonuçlar başarıyla 'output.json' dosyasına kaydedildi!")

      # Excel Dönüştürücü Fonksiyonu Çağırıyoruz:
      json_to_excel("output.json", "output.xlsx")

    except Exception as e:
      print(f"Hata oluştu: {e}")